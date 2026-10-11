"""One organized, printable workbook for each production shift report."""

from collections import defaultdict
from datetime import date, datetime
import math

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "17365D"
TEAL = "246B80"
INK = "233746"
MUTED = "607384"
PALE = "F0F5F8"
LINE = "DCE5EC"
NUMBER_FORMAT = "#,##0.##;[Red](#,##0.##);0"


def text_value(value):
    """Keep user-entered labels as text, even if they start with a formula."""
    value = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def local_datetime(value):
    if value is None:
        return "غير مسجل"
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.replace(tzinfo=None)


def write_cell(sheet, row, column, value, *, heading=False):
    if isinstance(value, str):
        value = text_value(value)
    elif isinstance(value, float) and not math.isfinite(value):
        value = "غير متاح"
    cell = sheet.cell(row, column, value)
    cell.font = Font(name="Arial", size=11, color="FFFFFF" if heading else INK, bold=heading)
    cell.alignment = Alignment(
        horizontal="center" if heading or isinstance(value, (int, float, datetime)) else "right",
        vertical="center" if heading else "top", wrap_text=True,
    )
    cell.border = Border(bottom=Side(style="thin", color=LINE))
    if heading:
        cell.fill = PatternFill("solid", fgColor=TEAL)
    elif row % 2 == 0:
        cell.fill = PatternFill("solid", fgColor=PALE)
    if isinstance(value, (int, float)):
        cell.number_format = NUMBER_FORMAT
    elif isinstance(value, datetime):
        cell.number_format = "yyyy-mm-dd hh:mm"
    elif isinstance(value, date):
        cell.number_format = "yyyy-mm-dd"
    return cell


def merged_text(sheet, row, first, last, value, *, title=False):
    sheet.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    cell = write_cell(sheet, row, first, value)
    cell.alignment = Alignment(horizontal="center" if title else "right", vertical="center", wrap_text=True)
    if title:
        cell.font = Font(name="Arial", size=18, color="FFFFFF", bold=True)
        for column in range(first, last + 1):
            sheet.cell(row, column).fill = PatternFill("solid", fgColor=NAVY)
        sheet.row_dimensions[row].height = 38
    return cell


def prepare_sheet(sheet, report, widths, *, landscape=False, a3=False):
    end = get_column_letter(len(widths))
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    merged_text(sheet, 1, 1, len(widths), "تقرير وردية الإنتاج", title=True)
    merged_text(sheet, 2, 1, len(widths), f"{report.factory.name} | {report.get_shift_display()} | {report.report_date:%Y-%m-%d}")
    merged_text(sheet, 3, 1, len(widths), f"رقم التقرير: {report.pk} | المشرف: {report.supervisor.display_name} | الحالة: {report.get_status_display()}")
    sheet.row_dimensions[2].height = 26
    sheet.row_dimensions[3].height = 26
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.45
    sheet.page_margins.bottom = 0.45
    sheet.oddFooter.right.text = "صفحة &P من &N"
    sheet.oddFooter.left.text = f"تقرير رقم {report.pk} | {report.report_date:%Y-%m-%d}"
    for footer in (sheet.oddFooter.left, sheet.oddFooter.right):
        footer.size = 9
        footer.font = "Arial"
        footer.color = MUTED
    sheet.print_area = f"A1:{end}3"


def write_table(sheet, header_row, headers, rows):
    for column, label in enumerate(headers, 1):
        write_cell(sheet, header_row, column, label, heading=True)
    sheet.row_dimensions[header_row].height = 40
    for row_index, values in enumerate(rows, header_row + 1):
        for column, value in enumerate(values, 1):
            write_cell(sheet, row_index, column, value)
        sheet.row_dimensions[row_index].height = 42
    end = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A{header_row}:{end}{max(header_row, sheet.max_row)}"
    sheet.freeze_panes = f"F{header_row + 1}" if len(headers) > 10 else f"A{header_row + 1}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.print_area = f"A1:{end}{sheet.max_row}"


def production_quantity(entry):
    # Entries created before unit-aware reporting may only have the kg column.
    if entry.final_production_quantity == 0 and entry.final_production_weight_kg > 0 and entry.final_production_unit_name in ("", "كجم"):
        return entry.final_production_weight_kg
    return entry.final_production_quantity


def build_report_workbook(report, entries, stoppages):
    entries = list(entries)
    stoppages = list(stoppages)
    workbook = Workbook()
    workbook.properties.title = f"تقرير إنتاج {report.factory.name} {report.report_date}"
    workbook.properties.subject = report.get_shift_display()
    workbook.properties.creator = "نظام الوسام"
    workbook.properties.language = "ar-EG"

    summary = workbook.active
    summary.title = "ملخص الوردية"
    prepare_sheet(summary, report, [23, 21, 21, 21, 20, 20])
    metadata = [
        ("تاريخ الوردية", report.report_date),
        ("بداية الوردية", local_datetime(report.started_at_device)),
        ("نهاية الوردية", local_datetime(report.completed_at_device)),
        ("المشرف المستلم", report.handover_to_supervisor.display_name if report.handover_to_supervisor else "غير محدد"),
        ("توقيت الاستلام", local_datetime(report.handover_confirmed_at)),
        ("عدد الماكينات المسجلة", len(entries)),
        ("عدد التوقفات المسجلة", len(stoppages)),
        ("مجموع مدد التوقفات المسجلة (دقيقة)", sum(stop.duration_minutes for stop in stoppages)),
        ("عدد تغييرات المنتج", sum(entry.product_changed for entry in entries)),
        ("عدد تغييرات العامل", sum(entry.operator_changed for entry in entries)),
    ]
    for row, (label, value) in enumerate(metadata, 5):
        merged_text(summary, row, 1, 2, label)
        cell = merged_text(summary, row, 3, 6, value)
        cell.font = Font(name="Arial", size=11, color=INK, bold=True)
        summary.row_dimensions[row].height = 28

    grouped = defaultdict(float)
    unknown_weight_entries = 0
    known_weight = 0.0
    for entry in entries:
        quantity = production_quantity(entry)
        grouped[entry.final_production_unit_name or "كجم"] += quantity
        if quantity > 0 and entry.final_production_weight_kg == 0:
            unknown_weight_entries += 1
        else:
            known_weight += entry.final_production_weight_kg
    row = 16
    merged_text(summary, row, 1, 6, "إجمالي الإنتاج حسب الوحدة")
    for column, label in ((1, "الوحدة"), (3, "الكمية")):
        merged_text(summary, row + 1, column, column + (1 if column == 1 else 3), label)
        summary.cell(row + 1, column).font = Font(name="Arial", size=11, color=INK, bold=True)
    row += 2
    if not grouped:
        merged_text(summary, row, 1, 6, "لا توجد بيانات إنتاج مسجلة.")
        row += 1
    for unit, quantity in sorted(grouped.items()):
        merged_text(summary, row, 1, 2, unit)
        merged_text(summary, row, 3, 6, quantity)
        summary.row_dimensions[row].height = 25
        row += 1
    merged_text(summary, row + 1, 1, 2, "الوزن المحسوب (كجم)")
    merged_text(summary, row + 1, 3, 6, known_weight)
    if unknown_weight_entries:
        merged_text(summary, row + 2, 1, 6, f"الوزن المحسوب جزئي: {unknown_weight_entries} ماكينة لها إنتاج بوحدة لا تتوفر لها معادلة تحويل إلى كجم.")
        summary.row_dimensions[row + 2].height = 34
    row += 4
    for label, value in (("ملاحظات الوردية", report.general_notes), ("ملاحظات الاستلام", report.handover_notes)):
        merged_text(summary, row, 1, 6, label)
        summary.cell(row, 1).font = Font(name="Arial", size=11, color=INK, bold=True)
        merged_text(summary, row + 1, 1, 6, value or "لا توجد ملاحظات")
        summary.row_dimensions[row + 1].height = min(180, max(45, 20 * (len(value or "") // 85 + (value or "").count("\n") + 1)))
        row += 3
    summary.freeze_panes = "A5"
    summary.print_area = f"A1:F{summary.max_row}"

    production = workbook.create_sheet("إنتاج الماكينات")
    prepare_sheet(production, report, [7, 16, 24, 23, 26, 18, 17, 20, 16, 23, 20, 42], landscape=True, a3=True)
    headers = [
        "م", "كود الماكينة", "الماكينة", "العامل", "المنتج", "الإنتاج الفعلي", "وحدة الإنتاج", "الوزن المحسوب (كجم)",
        "حالة التشغيل", "الخامة", "العبوة", "ملاحظات الماكينة",
    ]
    rows = []
    for index, entry in enumerate(entries, 1):
        quantity = production_quantity(entry)
        weight = entry.final_production_weight_kg if quantity == 0 or entry.final_production_weight_kg > 0 else None
        rows.append([
            index, entry.asset.asset_code, entry.asset.display_name, entry.display_operator, entry.display_product,
            quantity, entry.final_production_unit_name or "كجم", weight,
            entry.get_operation_mode_display(), entry.raw_material,
            entry.packaging_type, entry.notes,
        ])
    write_table(production, 5, headers, rows)
    for row, entry in enumerate(entries, 6):
        for column, changed, original in ((4, entry.operator_changed, entry.original_operator_name), (5, entry.product_changed, entry.original_product_name)):
            if changed:
                production.cell(row, column).fill = PatternFill("solid", fgColor="FFF0D5")
                production.cell(row, column).comment = Comment(f"القيمة الأصلية قبل التغيير: {original or 'غير محددة'}", "نظام الوسام")
    if not entries:
        merged_text(production, 6, 1, 12, "لم تسجل بيانات ماكينات في هذا التقرير.")
        production.print_area = "A1:L6"

    settings = workbook.create_sheet("إعدادات التشغيل")
    prepare_sheet(settings, report, [7, 16, 24, 16, 16, 19, 19, 19, 19, 20, 22], landscape=True)
    write_table(settings, 5, [
        "م", "كود الماكينة", "الماكينة", "اللقم الأصلية", "اللقم الحالية",
        "الدورة الأصلية (ث)", "الدورة الحالية (ث)", "التبريد الأصلي (ث)", "التبريد الحالي (ث)",
        "مستهدف الدورة", "وحدة المستهدف",
    ], [
        [index, entry.asset.asset_code, entry.asset.display_name, entry.original_cavities, entry.current_cavities,
         entry.original_cycle_time_seconds, entry.cycle_time_seconds, entry.original_cooling_time_seconds, entry.cooling_time_seconds,
         entry.target_cycle_production, entry.target_cycle_unit_name]
        for index, entry in enumerate(entries, 1)
    ])
    for row, entry in enumerate(entries, 6):
        for original_col, current_col in ((4, 5), (6, 7), (8, 9)):
            if settings.cell(row, original_col).value != settings.cell(row, current_col).value:
                settings.cell(row, current_col).fill = PatternFill("solid", fgColor="FFF0D5")
    if not entries:
        merged_text(settings, 6, 1, 11, "لم تسجل إعدادات ماكينات في هذا التقرير.")
        settings.print_area = "A1:K6"

    stops = workbook.create_sheet("التوقفات والأعطال")
    prepare_sheet(stops, report, [7, 18, 26, 28, 48, 18, 48, 23], landscape=True)
    write_table(stops, 5, ["م", "كود الماكينة", "الماكينة / نطاق التوقف", "نوع التوقف", "وصف التوقف", "المدة (دقيقة)", "الإجراء المتخذ", "توقيت التسجيل"], [
        [index, stop.asset.asset_code if stop.asset else "—", stop.asset.display_name if stop.asset else "جميع ماكينات المصنع",
         stop.get_stoppage_type_display(), stop.description, stop.duration_minutes, stop.action_taken, local_datetime(stop.created_at)]
        for index, stop in enumerate(stoppages, 1)
    ])
    if not stoppages:
        merged_text(stops, 6, 1, 8, "لا توجد توقفات مسجلة في هذه الوردية.")
    else:
        total_row = stops.max_row + 2
        merged_text(stops, total_row, 1, 5, "مجموع مدد التوقفات المسجلة (دقيقة)")
        write_cell(stops, total_row, 6, sum(stop.duration_minutes for stop in stoppages))
        merged_text(stops, total_row + 1, 1, 8, "قد تتداخل التوقفات؛ مجموع مددها لا يمثل مدة توقف المصنع الفعلية.")
    stops.print_area = f"A1:H{stops.max_row}"
    workbook.active = 0
    return workbook
