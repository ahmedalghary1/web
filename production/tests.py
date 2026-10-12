import uuid
from io import BytesIO
from importlib import import_module
from types import SimpleNamespace
from django.apps import apps
from django.db import connection
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient
from rest_framework import status

from accounts.models import User
from factories.models import Factory
from assets.models import Asset
from production.models import (
    Product,
    MachineOperator,
    MachineProductionDefault,
    ProductionOption,
    ProductionShiftReport,
    MachineProductionEntry,
    ProductionStoppage,
)

class ProductionAPITestCase(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.factory, _ = Factory.objects.get_or_create(code="F1", defaults={"name": "مصنع النقاء للبلاستيك"})
        
        self.supervisor1 = User.objects.create_user(
            phone="01111111111",
            password="password123",
            name="م. أحمد مشرف",
            role=User.Role.PRODUCTION_SUPERVISOR,
            factory=self.factory,
        )
        self.supervisor2 = User.objects.create_user(
            phone="01111111112",
            password="password123",
            name="م. محمود مستلم",
            role=User.Role.PRODUCTION_SUPERVISOR,
            factory=self.factory,
        )


        self.asset1 = Asset.objects.create(
            factory=self.factory,
            asset_code="M-01",
            asset_type=Asset.Type.REGULAR_MACHINE,
            sequence_order=1,
        )

        self.prod_default = Product.objects.create(
            factory=self.factory,
            name="برطمان 500 مل",
            code="PRD-500",
            weight_per_piece_grams=25.0,
        )
        self.prod_alternative = Product.objects.create(
            factory=self.factory,
            name="غطاء برطمان",
            code="CAP-500",
            weight_per_piece_grams=5.0,
        )

        self.operator = MachineOperator.objects.create(
            factory=self.factory,
            name="علي حسن",
            phone="01011112222",
        )

        self.machine_default = MachineProductionDefault.objects.create(
            asset=self.asset1,
            default_product=self.prod_default,
            original_cavities=4,
            cooling_time_seconds=12.5,
            cycle_time_seconds=22.0,
            target_cycle_production=650.0,
        )

    def test_table_snapshot_and_free_text_faults_survive_resync(self):
        import json
        self.client.force_authenticate(user=self.supervisor1)
        defaults = {
            "operator_name": "", "product_name": self.prod_default.name,
            "original_cavities": "4", "current_cavities": "4",
            "cooling_time_seconds": "12.5", "cycle_time_seconds": "22.0",
            "raw_material": "خامة افتراضية", "packaging_type": "كرتون",
            "final_production_unit_name": "كجم", "operation_mode": "AUTO",
            "target_cycle_production": "650.0", "target_cycle_unit_name": "",
        }
        payload = {
            "client_report_id": str(uuid.uuid4()), "report_date": timezone.localdate().isoformat(),
            "shift": "FIRST", "status": "DRAFT",
            "entries": [{"asset_id": self.asset1.pk, "product_name": "منتج مكتوب يدويًا",
                         "operator_name": "عامل جديد", "raw_material": "", "packaging_type": "",
                         "final_production_weight_kg": 12,
                         "dashboard_defaults": json.dumps(defaults, ensure_ascii=False)}],
            "stoppages": [{"stoppage_type": "OTHER", "description": "عطل أول\nعطل ثان", "duration_minutes": 0}],
        }
        url = "/api/v1/production/sync/shift-reports/"
        response = self.client.post(url, payload, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        entry = MachineProductionEntry.objects.get()
        self.assertEqual(entry.raw_material, "")
        self.assertEqual(entry.packaging_type, "")
        self.assertTrue(entry.product_changed)
        self.assertEqual(entry.original_product_name, self.prod_default.name)
        self.machine_default.default_product = self.prod_alternative
        self.machine_default.save()
        payload["entries"][0]["product_name"] = self.prod_default.name
        response = self.client.post(url, payload, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        entry = MachineProductionEntry.objects.get()
        self.assertFalse(entry.product_changed)
        self.assertEqual(json.loads(entry.dashboard_defaults), defaults)
        self.assertEqual(ProductionStoppage.objects.count(), 1)
        self.assertEqual(ProductionStoppage.objects.get().description, "عطل أول\nعطل ثان")
        self.assertTrue(any(row["changed"] for row in entry.default_comparison))
        self.client.force_login(self.supervisor1)
        report = self.client.get(reverse("production:report_detail", args=[entry.report_id]))
        self.assertEqual(report.status_code, 200)
        self.assertContains(report, "الافتراضي وقت التسجيل")
        self.assertContains(report, "عامل جديد")

    def test_table_rejects_invalid_snapshot(self):
        from production.serializers import SyncMachineEntryInputSerializer
        for value in ["[]", "null", '{"operator_name": []}', "invalid"]:
            serializer = SyncMachineEntryInputSerializer(data={"asset_id": self.asset1.pk, "dashboard_defaults": value})
            self.assertFalse(serializer.is_valid(), value)

    def test_bootstrap_endpoint(self):
        self.client.force_authenticate(user=self.supervisor1)
        res = self.client.get("/api/v1/production/bootstrap/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        self.assertEqual(len(data["assets"]), 1)
        self.assertEqual(len(data["products"]), 2)
        self.assertEqual(len(data["operators"]), 1)
        # Check machine default populated
        asset_data = data["assets"][0]
        self.assertIsNotNone(asset_data["production_default"])
        self.assertEqual(asset_data["production_default"]["original_cavities"], 4)
        self.assertEqual(asset_data["production_default"]["cooling_time_seconds"], 12.5)

    def test_sync_shift_report_with_product_change(self):
        self.client.force_authenticate(user=self.supervisor1)
        client_id = str(uuid.uuid4())
        today = str(timezone.localdate())

        payload = {
            "client_report_id": client_id,
            "shift": "FIRST",
            "report_date": today,
            "status": "PENDING_HANDOVER",
            "general_notes": "وردية منتظمة",
            "completed_at_device": timezone.now().isoformat(),
            "entries": [
                {
                    "asset_id": self.asset1.id,
                    "operator_id": self.operator.id,
                    "operator_name": "علي حسن",
                    "product_id": self.prod_alternative.id, # Changed from PRD-500!
                    "product_name": "غطاء برطمان",
                    "original_cavities": 4,
                    "current_cavities": 3, # One cavity blocked
                    "operation_mode": "AUTO",
                    "cooling_time_seconds": 12.5,
                    "cycle_time_seconds": 22.0,
                    "raw_material": "PP سابك 500P",
                    "final_production_weight_kg": 450.5,
                    "target_cycle_production": 650.0,
                    "packaging_type": "شكاير",
                    "notes": "تم تعطيل لقمة واحدة لوجود رايش",
                }
            ],
            "stoppages": [
                {
                    "asset_id": None, # General
                    "stoppage_type": "FRIDAY_PRAYER",
                    "description": "صلاة الجمعة لجميع الماكينات",
                    "duration_minutes": 60,
                    "action_taken": "إيقاف جماعي مؤقت واستئناف بعد الصلاة",
                }
            ],
        }

        res = self.client.post("/api/v1/production/sync/shift-reports/", payload, format="json")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        self.assertEqual(data["client_report_id"], client_id)
        self.assertEqual(data["status"], "PENDING_HANDOVER")
        self.assertEqual(len(data["machine_entries"]), 1)
        entry = data["machine_entries"][0]
        # Must detect product change!
        self.assertTrue(entry["product_changed"])
        self.assertEqual(entry["current_cavities"], 3)
        self.assertEqual(entry["original_cavities"], 4)


        # Stoppage
        self.assertEqual(len(data["stoppages"]), 1)
        self.assertEqual(data["stoppages"][0]["stoppage_type"], "FRIDAY_PRAYER")

        # Now test confirm handover by supervisor2
        self.client.force_authenticate(user=self.supervisor2)
        handover_res = self.client.post("/api/v1/production/reports/confirm-handover/", {
            "client_report_id": client_id,
            "handover_notes": "تم فحص الماكينات واستلام الوردية بالكامل بحالة جيدة",
        }, format="json")
        self.assertEqual(handover_res.status_code, status.HTTP_200_OK)
        handover_data = handover_res.json()
        self.assertEqual(handover_data["status"], "CONFIRMED")
        self.assertEqual(handover_data["handover_to_supervisor_name"], "م. محمود مستلم")

    def test_web_views_render(self):
        from django.test import Client
        web_client = Client()
        web_client.force_login(self.supervisor1)

        urls = [
            reverse("production:home"),
            reverse("production:reports"),
            reverse("production:machine_defaults"),
            reverse("production:products"),
            reverse("production:operators"),
            reverse("production:stoppages"),
        ]
        for url in urls:
            response = web_client.get(url)
            self.assertEqual(response.status_code, 200, f"URL {url} failed with {response.status_code}")

    def test_admin_can_access_without_factory(self):
        """Verify superuser / admin with factory=None has full access and never receives factory error."""
        admin_user = User.objects.create_superuser(
            phone="01099999999",
            password="adminpassword123",
            name="المدير العام",
        )
        self.assertIsNone(admin_user.factory)
        self.assertTrue(admin_user.is_admin)

        from django.test import Client
        admin_client = Client()
        admin_client.force_login(admin_user)

        # 1. Access production home
        response = admin_client.get(reverse("production:home"))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "المستخدم غير مربوط بمصنع.")

        # 2. Access with factory switcher param
        response_all = admin_client.get(reverse("production:home") + "?factory=all")
        self.assertEqual(response_all.status_code, 200)
        self.assertNotContains(response_all, "المستخدم غير مربوط بمصنع.")

        response_f1 = admin_client.get(reverse("production:home") + f"?factory={self.factory.id}")
        self.assertEqual(response_f1.status_code, 200)

        # 3. Access other production views
        self.assertEqual(admin_client.get(reverse("production:reports")).status_code, 200)
        self.assertEqual(admin_client.get(reverse("production:machine_defaults")).status_code, 200)
        self.assertEqual(admin_client.get(reverse("production:products")).status_code, 200)
        self.assertEqual(admin_client.get(reverse("production:operators")).status_code, 200)
        self.assertEqual(admin_client.get(reverse("production:stoppages")).status_code, 200)

    def test_maintenance_supervisor_cannot_access_production(self):
        """User with only maintenance permissions is redirected from production."""
        maint_supervisor = User.objects.create_user(
            phone="01200000001",
            password="password123",
            name="مشرف صيانة فقط",
            role=User.Role.MAINTENANCE_SUPERVISOR,
            factory=self.factory,
        )
        self.assertTrue(maint_supervisor.can_access_maintenance)
        self.assertFalse(maint_supervisor.can_access_production)

        from django.test import Client
        client = Client()
        client.force_login(maint_supervisor)

        response = client.get(reverse("production:home"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("dashboard:home"))

    def test_factory_isolation_between_supervisors(self):
        """Supervisor in Factory 1 cannot see reports from Factory 2."""
        factory2, _ = Factory.objects.get_or_create(code="F2", defaults={"name": "مصنع الوسام 2"})
        supervisor_f2 = User.objects.create_user(
            phone="01200000002",
            password="password123",
            name="مشرف مصنع 2",
            role=User.Role.PRODUCTION_SUPERVISOR,
            factory=factory2,
        )

        # Create report in Factory 2
        report_f2 = ProductionShiftReport.objects.create(
            client_report_id=str(uuid.uuid4()),
            factory=factory2,
            supervisor=supervisor_f2,
            shift="FIRST",
            report_date=timezone.localdate(),
        )

        # Create report in Factory 1
        report_f1 = ProductionShiftReport.objects.create(
            client_report_id=str(uuid.uuid4()),
            factory=self.factory,
            supervisor=self.supervisor1,
            shift="FIRST",
            report_date=timezone.localdate(),
        )

        from django.test import Client
        client1 = Client()
        client1.force_login(self.supervisor1)

        # Supervisor 1 views reports
        res = client1.get(reverse("production:reports"))
        self.assertEqual(res.status_code, 200)
        reports_in_view = res.context["page_obj"].object_list
        self.assertIn(report_f1, reports_in_view)
        self.assertNotIn(report_f2, reports_in_view)

        # Supervisor 1 cannot access detail of report in Factory 2
        detail_res = client1.get(reverse("production:report_detail", args=[report_f2.id]))
        self.assertEqual(detail_res.status_code, 404)

    def test_report_detail_displays_changed_values_clearly(self):
        from django.test import Client
        report = ProductionShiftReport.objects.create(
            client_report_id=str(uuid.uuid4()),
            factory=self.factory,
            supervisor=self.supervisor1,
            shift="FIRST",
            report_date=timezone.localdate(),
        )
        MachineProductionEntry.objects.create(
            report=report,
            asset=self.asset1,
            operator_name="أحمد محمود",
            original_operator_name="محمد علي",
            operator_changed=True,
            product_name="جركن 20 لتر جديد",
            original_product_name="جركن 10 لتر قديم",
            product_changed=True,
            final_production_weight_kg=150.0,
        )

        client = Client()
        client.force_login(self.supervisor1)

        # 1. Check report detail page displays explicit change text
        res = client.get(reverse("production:report_detail", args=[report.id]))
        self.assertEqual(res.status_code, 200)
        content = res.content.decode("utf-8")
        self.assertIn("تم تغيير العامل!", content)
        self.assertIn("كانت القيمة:", content)
        self.assertIn("محمد علي", content)
        self.assertIn("أحمد محمود", content)
        self.assertIn("تم تغيير المنتج!", content)
        self.assertIn("جركن 10 لتر قديم", content)
        self.assertIn("جركن 20 لتر جديد", content)
        self.assertIn("وتم تغييرها إلى:", content)

        # 2. Check home dashboard reflects changes in stats and recent reports
        home_res = client.get(reverse("production:home"))
        self.assertEqual(home_res.status_code, 200)
        home_content = home_res.content.decode("utf-8")
        self.assertIn("تغييرات العمال والمنتجات اليوم", home_content)
        self.assertIn("تغيير منتج", home_content)
        self.assertIn("تغيير عامل", home_content)

    def test_machine_defaults_view_and_edit_use_asset_code(self):
        from django.test import Client
        client = Client()
        client.force_login(self.supervisor1)

        # 1. Machine defaults list renders asset_code and does not have machine name column
        res = client.get(reverse("production:machine_defaults"))
        self.assertEqual(res.status_code, 200)
        content = res.content.decode("utf-8")
        self.assertIn("M-01", content)
        self.assertNotIn("اسم الماكينة", content)

        # 2. Machine default edit page renders asset_code
        edit_url = reverse("production:machine_default_edit", args=[self.asset1.id])
        edit_res = client.get(edit_url)
        self.assertEqual(edit_res.status_code, 200)
        edit_content = edit_res.content.decode("utf-8")
        self.assertIn("M-01", edit_content)

        # 3. Post edit form and check redirect and success message with asset_code
        post_res = client.post(edit_url, {
            "default_product": self.prod_default.id,
            "default_operator": self.operator.id,
            "original_cavities": 4,
            "cooling_time_seconds": 12.5,
            "cycle_time_seconds": 25.0,
            "target_cycle_production": 400.0,
        }, follow=True)
        self.assertEqual(post_res.status_code, 200)
        post_content = post_res.content.decode("utf-8")
        self.assertIn("تم تحديث إعدادات الإنتاج لماكينة M-01 بنجاح.", post_content)
        self.assertIn("M-01", post_content)

    def test_three_shifts_automatic_handover_logic(self):
        # 1. Create 3 supervisors assigned to the 3 shifts in factory 1
        sup1 = User.objects.create_user(
            phone="01011111111",
            password="pass",
            name="مشرف وردية 1",
            role=User.Role.PRODUCTION_SUPERVISOR,
            factory=self.factory,
            shift=User.Shift.FIRST
        )
        sup2 = User.objects.create_user(
            phone="01022222222",
            password="pass",
            name="مشرف وردية 2",
            role=User.Role.PRODUCTION_SUPERVISOR,
            factory=self.factory,
            shift=User.Shift.SECOND
        )
        sup3 = User.objects.create_user(
            phone="01033333333",
            password="pass",
            name="مشرف وردية 3",
            role=User.Role.PRODUCTION_SUPERVISOR,
            factory=self.factory,
            shift=User.Shift.THIRD
        )

        # Check cyclic next/previous shifts
        self.assertEqual(sup1.next_shift, "SECOND")
        self.assertEqual(sup1.get_next_shift_supervisor(), sup2)
        self.assertEqual(sup2.next_shift, "THIRD")
        self.assertEqual(sup2.get_next_shift_supervisor(), sup3)
        self.assertEqual(sup3.next_shift, "FIRST")
        self.assertEqual(sup3.get_next_shift_supervisor(), sup1)

        self.assertEqual(sup1.previous_shift, "THIRD")
        self.assertEqual(sup2.previous_shift, "FIRST")
        self.assertEqual(sup3.previous_shift, "SECOND")

        # 2. Supervisor 1 finishes FIRST shift report
        client_rep_id = str(uuid.uuid4())
        today_str = str(timezone.localdate())

        self.client.force_authenticate(user=sup1)
        sync_payload = {
            "client_report_id": client_rep_id,
            "shift": "FIRST",
            "report_date": today_str,
            "status": "PENDING_HANDOVER",
            "general_notes": "تم تسليم الوردية 1 بنجاح",
            "completed_at_device": timezone.now().isoformat(),
            "entries": [
                {
                    "asset_id": self.asset1.id,
                    "product_id": self.prod_default.id,
                    "original_cavities": 4,
                    "current_cavities": 4,
                    "final_production_weight_kg": 250.0,
                }
            ],
            "stoppages": []
        }
        res_sync = self.client.post(reverse("production-api:sync-shift-reports"), sync_payload, format="json")
        self.assertEqual(res_sync.status_code, 200)

        # Check report in DB: handover_to_supervisor should be automatically sup2!
        rep = ProductionShiftReport.objects.get(client_report_id=client_rep_id)
        self.assertEqual(rep.supervisor, sup1)
        self.assertEqual(rep.handover_to_supervisor, sup2)
        self.assertEqual(rep.status, ProductionShiftReport.Status.PENDING_HANDOVER)

        # 3. Supervisor 2 logs in / opens app -> calls Bootstrap
        self.client.force_authenticate(user=sup2)
        res_bootstrap = self.client.get(reverse("production-api:bootstrap"))
        self.assertEqual(res_bootstrap.status_code, 200)
        data = res_bootstrap.data

        # Bootstrap user info has shift and next supervisor (sup3)
        self.assertEqual(data["user"]["shift"], "SECOND")
        self.assertEqual(data["user"]["next_shift"], "THIRD")
        self.assertEqual(data["user"]["next_shift_supervisor"]["id"], sup3.id)
        self.assertEqual(data["user"]["next_shift_supervisor"]["name"], "مشرف وردية 3")

        # Pending handover is automatically sup1's report!
        self.assertIsNotNone(data["pending_handover"])
        self.assertEqual(data["pending_handover"]["client_report_id"], client_rep_id)
        self.assertEqual(data["pending_handover"]["shift"], "FIRST")

        # 4. Supervisor 2 confirms handover
        res_confirm = self.client.post(
            reverse("production-api:confirm-handover"),
            {"client_report_id": client_rep_id, "handover_notes": "تم استلام جميع الماكينات بحالة جيدة"},
            format="json"
        )
        self.assertEqual(res_confirm.status_code, 200)

        rep.refresh_from_db()
        self.assertEqual(rep.status, ProductionShiftReport.Status.CONFIRMED)
        self.assertEqual(rep.handover_to_supervisor, sup2)
        self.assertIsNotNone(rep.handover_confirmed_at)
        self.assertEqual(rep.handover_notes, "تم استلام جميع الماكينات بحالة جيدة")

        # Calling bootstrap again for sup2 shows no pending handover
        res_bootstrap2 = self.client.get(reverse("production-api:bootstrap"))
        self.assertIsNone(res_bootstrap2.data["pending_handover"])

    def test_empty_sync_never_wipes_recorded_machine_entries(self):
        client_rep_id = str(uuid.uuid4())
        today_str = str(timezone.localdate())

        self.client.force_authenticate(user=self.supervisor1)
        
        # 1. First sync sends 1 recorded machine entry
        initial_sync = {
            "client_report_id": client_rep_id,
            "shift": "FIRST",
            "report_date": today_str,
            "status": "PENDING_HANDOVER",
            "general_notes": "تسجيل ماكينة تجريبية",
            "entries": [
                {
                    "asset_id": self.asset1.id,
                    "product_id": self.prod_default.id,
                    "original_cavities": 4,
                    "current_cavities": 4,
                    "final_production_weight_kg": 320.5,
                }
            ],
            "stoppages": []
        }
        res1 = self.client.post(reverse("production-api:sync-shift-reports"), initial_sync, format="json")
        self.assertEqual(res1.status_code, 200)

        rep = ProductionShiftReport.objects.get(client_report_id=client_rep_id)
        self.assertEqual(rep.machine_entries.count(), 1)
        entry = rep.machine_entries.first()
        self.assertEqual(entry.asset, self.asset1)
        self.assertEqual(entry.final_production_weight_kg, 320.5)

        # 2. Subsequent sync sent with empty entries list (e.g. status update or finish shift without entries)
        finish_sync = {
            "client_report_id": client_rep_id,
            "shift": "FIRST",
            "report_date": today_str,
            "status": "PENDING_HANDOVER",
            "general_notes": "تم إنهاء الوردية",
            "entries": [],
            "stoppages": []
        }
        res2 = self.client.post(reverse("production-api:sync-shift-reports"), finish_sync, format="json")
        self.assertEqual(res2.status_code, 200)

        # Machine entries MUST NOT be wiped!
        rep.refresh_from_db()
        self.assertEqual(rep.machine_entries.count(), 1, "Recorded machine entries must be preserved when sync entries is empty")
        self.assertEqual(rep.general_notes, "تم إنهاء الوردية")

    def sync_payload(self, entry=None, **overrides):
        payload = {
            "client_report_id": str(uuid.uuid4()),
            "factory_id": self.factory.pk,
            "shift": "FIRST",
            "report_date": str(timezone.localdate()),
            "status": "DRAFT",
            "entries": [entry if entry is not None else {"asset_id": self.asset1.pk}],
        }
        payload.update(overrides)
        return payload

    def test_sync_preserves_user_originals_and_zero_actual_times(self):
        self.client.force_authenticate(self.supervisor1)
        payload = self.sync_payload({
            "asset_id": self.asset1.pk,
            "original_cavities": 6,
            "current_cavities": 5,
            "original_cycle_time_seconds": 22,
            "original_cooling_time_seconds": 12.5,
            "cycle_time_seconds": 0,
            "cooling_time_seconds": 0,
        })
        response = self.client.post(reverse("production-api:sync-shift-reports"), payload, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        entry = response.data["machine_entries"][0]
        self.assertEqual((entry["original_cavities"], entry["current_cavities"]), (6, 5))
        self.assertEqual((entry["original_cycle_time_seconds"], entry["cycle_time_seconds"]), (22, 0))
        self.assertEqual((entry["original_cooling_time_seconds"], entry["cooling_time_seconds"]), (12.5, 0))
        self.machine_default.refresh_from_db()
        self.assertEqual(self.machine_default.original_cavities, 4)
        self.assertEqual(self.machine_default.cycle_time_seconds, 22)

    def test_omitted_settings_prefill_defaults_then_keep_saved_snapshot_after_default_change(self):
        self.client.force_authenticate(self.supervisor1)
        payload = self.sync_payload()
        first = self.client.post(reverse("production-api:sync-shift-reports"), payload, format="json")
        self.assertEqual(first.status_code, 200)
        entry = first.data["machine_entries"][0]
        self.assertEqual((entry["original_cavities"], entry["current_cavities"]), (4, 4))
        self.assertEqual((entry["original_cycle_time_seconds"], entry["cycle_time_seconds"]), (22, 22))
        self.assertEqual((entry["original_cooling_time_seconds"], entry["cooling_time_seconds"]), (12.5, 12.5))
        self.machine_default.original_cavities = 10
        self.machine_default.cycle_time_seconds = 35
        self.machine_default.cooling_time_seconds = 20
        self.machine_default.save()
        payload["entries"][0].update({"cycle_time_seconds": 23, "original_cooling_time_seconds": None})
        second = self.client.post(reverse("production-api:sync-shift-reports"), payload, format="json")
        self.assertEqual(second.status_code, 200)
        entry = second.data["machine_entries"][0]
        self.assertEqual((entry["original_cavities"], entry["current_cavities"]), (4, 4))
        self.assertEqual((entry["original_cycle_time_seconds"], entry["cycle_time_seconds"]), (22, 23))
        self.assertEqual((entry["original_cooling_time_seconds"], entry["cooling_time_seconds"]), (12.5, 12.5))
        self.assertEqual(ProductionShiftReport.objects.count(), 1)
        self.assertEqual(MachineProductionEntry.objects.count(), 1)

    def test_sync_rejects_invalid_measurements_without_creating_report(self):
        self.client.force_authenticate(self.supervisor1)
        cases = [(field, value) for field in (
            "original_cycle_time_seconds", "original_cooling_time_seconds", "cycle_time_seconds", "cooling_time_seconds",
            "final_production_quantity", "final_production_weight_kg", "target_cycle_production",
        ) for value in (-1, "NaN", "Infinity")]
        cases.extend([(field, 0) for field in ("original_cavities", "current_cavities")])
        for field, value in cases:
            with self.subTest(field=field, value=value):
                payload = self.sync_payload({"asset_id": self.asset1.pk, field: value})
                response = self.client.post(reverse("production-api:sync-shift-reports"), payload, format="json")
                self.assertEqual(response.status_code, 400, response.data)
                self.assertEqual(ProductionShiftReport.objects.count(), 0)

    def test_sync_rejects_foreign_factory_report_and_stoppages(self):
        other = Factory.objects.create(name="مصنع آخر", code="OTHER")
        other_asset = Asset.objects.create(factory=other, asset_code="X-1", asset_type=Asset.Type.REGULAR_MACHINE)
        foreign = ProductionShiftReport.objects.create(factory=other, supervisor=self.supervisor2, report_date=timezone.localdate(), status="DRAFT")
        self.client.force_authenticate(self.supervisor1)
        for overrides in (
            {"factory_id": other.pk},
            {"client_report_id": str(foreign.client_report_id)},
            {"stoppages": [{"asset_id": other_asset.pk, "description": "توقف", "duration_minutes": 5}]},
        ):
            with self.subTest(overrides=overrides):
                response = self.client.post(reverse("production-api:sync-shift-reports"), self.sync_payload(**overrides), format="json")
                self.assertIn(response.status_code, (400, 403))
                self.assertEqual(ProductionShiftReport.objects.count(), 1)
        foreign.refresh_from_db()
        self.assertEqual(foreign.status, "DRAFT")

    def test_admin_sync_uses_explicit_factory_id(self):
        admin = User.objects.create_superuser(phone="01199998888", password="pass")
        other = Factory.objects.create(name="مصنع الإدارة", code="ADM")
        other_asset = Asset.objects.create(factory=other, asset_code="A-1", asset_type=Asset.Type.REGULAR_MACHINE)
        self.client.force_authenticate(admin)
        response = self.client.post(reverse("production-api:sync-shift-reports"), self.sync_payload(
            {"asset_id": other_asset.pk}, factory_id=other.pk,
        ), format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["factory"], other.pk)

    def test_confirmed_report_cannot_be_rewritten_by_sync(self):
        report = ProductionShiftReport.objects.create(
            factory=self.factory, supervisor=self.supervisor1, report_date=timezone.localdate(),
            status="CONFIRMED", completed_at_device=timezone.now(), general_notes="معتمد",
        )
        entry = MachineProductionEntry.objects.create(report=report, asset=self.asset1, current_cavities=3, cycle_time_seconds=27)
        self.client.force_authenticate(self.supervisor1)
        response = self.client.post(reverse("production-api:sync-shift-reports"), self.sync_payload(
            {"asset_id": self.asset1.pk, "current_cavities": 10, "cycle_time_seconds": 99},
            client_report_id=str(report.client_report_id), general_notes="تعديل",
        ), format="json")
        self.assertEqual(response.status_code, 200)
        report.refresh_from_db()
        entry.refresh_from_db()
        self.assertEqual(report.status, "CONFIRMED")
        self.assertEqual(report.general_notes, "معتمد")
        self.assertEqual((entry.current_cavities, entry.cycle_time_seconds), (3, 27))

    def test_different_client_uuid_cannot_take_over_existing_shift(self):
        report = ProductionShiftReport.objects.create(
            factory=self.factory, supervisor=self.supervisor1, report_date=timezone.localdate(), status="DRAFT", general_notes="تقرير أصلي",
        )
        entry = MachineProductionEntry.objects.create(report=report, asset=self.asset1, current_cavities=3)
        original_id = report.client_report_id
        self.client.force_authenticate(self.supervisor2)
        response = self.client.post(reverse("production-api:sync-shift-reports"), self.sync_payload(
            {"asset_id": self.asset1.pk, "current_cavities": 9}, general_notes="تقرير من هاتف آخر",
        ), format="json")
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.data["client_report_id"], str(original_id))
        report.refresh_from_db()
        entry.refresh_from_db()
        self.assertEqual(report.client_report_id, original_id)
        self.assertEqual(report.general_notes, "تقرير أصلي")
        self.assertEqual(entry.current_cavities, 3)

    def test_sync_retains_inactive_choices_from_offline_draft(self):
        self.prod_default.is_active = False
        self.prod_default.save()
        self.operator.is_active = False
        self.operator.save()
        unit = ProductionOption.objects.create(factory=self.factory, category="FINAL_UNIT", name="وحدة قديمة", is_active=False)
        raw = ProductionOption.objects.create(factory=self.factory, category="RAW_MATERIAL", name="خامة قديمة", is_active=False)
        pack = ProductionOption.objects.create(factory=self.factory, category="PACKAGING", name="عبوة قديمة", is_active=False)
        self.client.force_authenticate(self.supervisor1)
        payload = self.sync_payload({
            "asset_id": self.asset1.pk, "product_id": self.prod_default.pk, "operator_id": self.operator.pk,
            "final_production_unit_id": unit.pk, "raw_material_option_id": raw.pk, "packaging_option_id": pack.pk,
            "final_production_quantity": 10,
        })
        first = self.client.post(reverse("production-api:sync-shift-reports"), payload, format="json")
        self.assertEqual(first.status_code, 200, first.data)
        entry = first.data["machine_entries"][0]
        self.assertEqual((entry["product"], entry["operator"], entry["final_production_unit"]), (self.prod_default.pk, self.operator.pk, unit.pk))
        self.assertEqual((entry["raw_material"], entry["packaging_type"]), (raw.name, pack.name))
        for field in ("raw_material_option_id", "packaging_option_id", "final_production_unit_id"):
            payload["entries"][0].pop(field)
        second = self.client.post(reverse("production-api:sync-shift-reports"), payload, format="json")
        self.assertEqual(second.status_code, 200, second.data)
        entry = second.data["machine_entries"][0]
        self.assertEqual((entry["raw_material"], entry["packaging_type"], entry["final_production_unit_name"]), (raw.name, pack.name, unit.name))

    def test_unfinished_legacy_report_cannot_be_handed_over(self):
        report = ProductionShiftReport.objects.create(
            factory=self.factory, supervisor=self.supervisor1, report_date=timezone.localdate(), status="PENDING_HANDOVER",
        )
        self.client.force_authenticate(self.supervisor2)
        bootstrap = self.client.get(reverse("production-api:bootstrap"))
        self.assertIsNone(bootstrap.data["pending_handover"])
        result = self.client.post(reverse("production-api:confirm-handover"), {"client_report_id": str(report.client_report_id)}, format="json")
        self.assertEqual(result.status_code, 400)
        report.refresh_from_db()
        self.assertEqual(report.status, "PENDING_HANDOVER")

    def test_historical_timing_backfill_uses_recorded_report_values(self):
        report = ProductionShiftReport.objects.create(factory=self.factory, supervisor=self.supervisor1, report_date=timezone.localdate())
        entry = MachineProductionEntry.objects.create(report=report, asset=self.asset1, cycle_time_seconds=17.5, cooling_time_seconds=8)
        migration = import_module("production.migrations.0004_report_timing_snapshots")
        migration.preserve_recorded_times(apps, SimpleNamespace(connection=connection))
        entry.refresh_from_db()
        self.assertEqual(entry.original_cycle_time_seconds, 17.5)
        self.assertEqual(entry.original_cooling_time_seconds, 8)
        self.assertNotEqual(entry.original_cycle_time_seconds, self.machine_default.cycle_time_seconds)

    def test_excel_report_contains_numeric_totals_snapshot_times_and_stoppages(self):
        from openpyxl import load_workbook
        report = ProductionShiftReport.objects.create(
            factory=self.factory, supervisor=self.supervisor1, report_date=timezone.localdate(),
            started_at_device=timezone.now(), completed_at_device=timezone.now(),
            general_notes='=HYPERLINK("https://example.invalid")', handover_notes="مراجعة المخزن",
        )
        MachineProductionEntry.objects.create(
            report=report, asset=self.asset1, operator_name="عامل جديد", original_operator_name="عامل سابق", operator_changed=True,
            product_name="+قيمة نصية", original_cavities=4, current_cavities=3,
            original_cycle_time_seconds=22, cycle_time_seconds=24, original_cooling_time_seconds=12.5, cooling_time_seconds=13,
            final_production_quantity=20, final_production_unit_name="كرتونة", notes="ملاحظات الإنتاج",
        )
        second_asset = Asset.objects.create(factory=self.factory, asset_code="M-02", asset_type=Asset.Type.REGULAR_MACHINE)
        MachineProductionEntry.objects.create(
            report=report, asset=second_asset, final_production_quantity=5.5, final_production_unit_name="كجم", final_production_weight_kg=5.5,
        )
        ProductionStoppage.objects.create(report=report, asset=self.asset1, description="تعطل الحساس", duration_minutes=15, action_taken="استبدال الحساس")
        self.client.force_login(self.supervisor1)
        response = self.client.get(reverse("production:report_export_excel", args=[report.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn(f"production-report-{report.pk}-", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["ملخص الوردية", "إنتاج الماكينات", "إعدادات التشغيل", "التوقفات والأعطال"])
        for sheet in workbook:
            self.assertTrue(sheet.sheet_view.rightToLeft)
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertTrue(sheet.freeze_panes)
            self.assertTrue(sheet.print_area)
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.oddFooter.right.text, "صفحة &P من &N")
            self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))
        production = workbook["إنتاج الماكينات"]
        self.assertEqual(production["F6"].value, 20)
        self.assertEqual(production["F6"].data_type, "n")
        self.assertEqual(production["G6"].value, "كرتونة")
        self.assertIsNone(production["H6"].value)
        self.assertIn("عامل سابق", production["D6"].comment.text)
        self.assertEqual(production["E6"].value, "'+قيمة نصية")
        settings = workbook["إعدادات التشغيل"]
        self.assertEqual([settings.cell(6, col).value for col in range(4, 10)], [4, 3, 22, 24, 12.5, 13])
        summary = workbook["ملخص الوردية"]
        all_values = [cell.value for row in summary for cell in row]
        self.assertIn("كرتونة", all_values)
        self.assertIn(20, all_values)
        self.assertIn(5.5, all_values)
        self.assertNotIn(25.5, all_values)
        self.assertTrue(any("الوزن المحسوب جزئي" in str(value) for value in all_values))
        stops = workbook["التوقفات والأعطال"]
        self.assertEqual((stops["E6"].value, stops["F6"].value, stops["G6"].value), ("تعطل الحساس", 15, "استبدال الحساس"))

    def test_excel_export_denies_another_factory_and_handles_empty_report(self):
        from openpyxl import load_workbook
        other = Factory.objects.create(name="مصنع آخر", code="EX2")
        foreign = ProductionShiftReport.objects.create(factory=other, supervisor=self.supervisor2, report_date=timezone.localdate())
        self.client.force_login(self.supervisor1)
        self.assertEqual(self.client.get(reverse("production:report_export_excel", args=[foreign.pk])).status_code, 404)
        report = ProductionShiftReport.objects.create(factory=self.factory, supervisor=self.supervisor1, report_date=timezone.localdate())
        response = self.client.get(reverse("production:report_export_excel", args=[report.pk]))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertIn("لم تسجل بيانات ماكينات", workbook["إنتاج الماكينات"]["A6"].value)
        self.assertIn("لا توجد توقفات", workbook["التوقفات والأعطال"]["A6"].value)





